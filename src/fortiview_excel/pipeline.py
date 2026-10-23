"""Normalize synthetic FortiView traffic samples and update an Excel workbook."""

from __future__ import annotations

import argparse
import re
import shutil
import os
import tempfile
import uuid
from contextlib import contextmanager
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

UNIT_FACTORS = {
    "B": 1,
    "KB": 1024,
    "MB": 1024**2,
    "GB": 1024**3,
    "TB": 1024**4,
}
HEADERS = ["collected_at", "device", "interface", "direction", "original_value", "bytes"]


def parse_traffic_size(value: str) -> int:
    match = re.fullmatch(r"\s*([0-9]+(?:[.,][0-9]+)?)\s*([KMGT]?B)\s*", str(value), re.I)
    if not match:
        raise ValueError(f"Unsupported traffic value: {value}")
    number = float(match.group(1).replace(",", "."))
    unit = match.group(2).upper()
    return round(number * UNIT_FACTORS[unit])


def load_traffic(path: str | Path) -> pd.DataFrame:
    frame = pd.read_csv(path)
    required = {"collected_at", "device", "interface", "direction", "value"}
    missing = required.difference(frame.columns)
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(sorted(missing))}")
    frame = frame.copy()
    frame["collected_at"] = pd.to_datetime(frame["collected_at"], errors="raise", utc=True)
    frame["bytes"] = frame["value"].map(parse_traffic_size)
    frame["direction"] = frame["direction"].str.casefold()
    return frame.sort_values(["collected_at", "device", "interface", "direction"])


def backup_workbook(path: Path) -> Path | None:
    if not path.exists():
        return None
    backup_dir = path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    backup = backup_dir / f"{path.stem}_{timestamp}_{uuid.uuid4().hex[:8]}{path.suffix}"
    shutil.copy2(path, backup)
    return backup


def _create_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F5F")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F1"
    return workbook


def _existing_keys(sheet) -> dict[tuple[str, str, str, str], int]:
    keys = {}
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is None:
            continue
        key = (str(row[0]), str(row[1]), str(row[2]), str(row[3]))
        if key in keys:
            raise ValueError('Existing workbook contains duplicate keys; reconcile before updating.')
        keys[key] = index
    return keys


def _copy_previous_style(sheet, target_row: int) -> None:
    if target_row <= 2:
        return
    for column in range(1, len(HEADERS) + 1):
        source = sheet.cell(target_row - 1, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format


def _refresh_summary(workbook: Workbook) -> None:
    if "Resumo" in workbook.sheetnames:
        del workbook["Resumo"]
    summary = workbook.create_sheet("Resumo", 0)
    summary.append(["device", "direction", "total_bytes", "total_gib"])
    data = workbook["Dados"]
    rows = list(data.iter_rows(min_row=2, values_only=True))
    frame = pd.DataFrame(rows, columns=HEADERS)
    if not frame.empty:
        grouped = (
            frame.groupby(["device", "direction"], as_index=False)["bytes"]
            .sum()
            .sort_values(["device", "direction"])
        )
        grouped["total_gib"] = grouped["bytes"] / 1024**3
        for row in grouped.itertuples(index=False):
            summary.append([row.device, row.direction, row.bytes, row.total_gib])
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16845C")
    summary.freeze_panes = "A2"


@contextmanager
def workbook_lock(output: Path):
    lock = output.with_suffix(output.suffix + '.lock')
    try:
        handle = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise RuntimeError('Another run owns the workbook lock. Check the running job before removing a stale lock.') from exc
    try:
        with os.fdopen(handle, 'w') as stream:
            stream.write(str(os.getpid()))
        yield
    finally:
        lock.unlink(missing_ok=True)


def _atomic_save(workbook, output):
    handle, name = tempfile.mkstemp(prefix=output.stem + '-', suffix='.xlsx', dir=output.parent)
    os.close(handle)
    temporary = Path(name)
    try:
        workbook.save(temporary)
        with temporary.open('r+b') as stream:
            os.fsync(stream.fileno())
        import zipfile
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip() is not None:
                raise ValueError('Invalid generated workbook')
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def update_workbook(input_path: str | Path, output_path: str | Path) -> dict[str, object]:
    frame = load_traffic(input_path).drop_duplicates()
    if frame.duplicated(['collected_at', 'device', 'interface', 'direction']).any():
        raise ValueError('Conflicting input values for the same traffic sample.')
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with workbook_lock(output):
        backup = backup_workbook(output)
        workbook = load_workbook(output) if output.exists() else _create_workbook()
        inserted = updated = unchanged = 0
        try:
            sheet = workbook['Dados']
            keys = _existing_keys(sheet)
            for row in frame.itertuples(index=False):
                timestamp = row.collected_at.isoformat()
                key = (timestamp, row.device, row.interface, row.direction)
                if key in keys:
                    index = keys[key]
                    if sheet.cell(index, 6).value == row.bytes:
                        unchanged += 1
                        continue
                    if 'Alteracoes' not in workbook.sheetnames:
                        audit = workbook.create_sheet('Alteracoes')
                        audit.append(['changed_at_utc', *HEADERS[:4], 'old_value', 'old_bytes', 'new_value', 'new_bytes'])
                    workbook['Alteracoes'].append([datetime.now(timezone.utc).isoformat(), *key,
                        sheet.cell(index, 5).value, sheet.cell(index, 6).value, row.value, row.bytes])
                    sheet.cell(index, 5, row.value)
                    sheet.cell(index, 6, row.bytes)
                    updated += 1
                    continue
                sheet.append([timestamp, row.device, row.interface, row.direction, row.value, row.bytes])
                _copy_previous_style(sheet, sheet.max_row)
                keys[key] = sheet.max_row
                inserted += 1
            for column, width in {'A': 28, 'B': 18, 'C': 16, 'D': 14, 'E': 16, 'F': 18}.items():
                sheet.column_dimensions[column].width = width
            sheet.auto_filter.ref = sheet.dimensions
            _refresh_summary(workbook)
            _atomic_save(workbook, output)
        finally:
            workbook.close()
    return {'output': output, 'backup': backup, 'inserted': inserted, 'updated': updated, 'unchanged': unchanged}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="data/traffic_sample.csv")
    parser.add_argument("--output", default="output/traffic_weekly.xlsx")
    args = parser.parse_args()
    result = update_workbook(args.input, args.output)
    print(f"{result['output']} | inserted={result['inserted']} | backup={result['backup']}")


if __name__ == "__main__":
    main()
