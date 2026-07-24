from pathlib import Path

from openpyxl import load_workbook

from fortiview_excel.pipeline import parse_traffic_size, update_workbook


ROOT = Path(__file__).resolve().parents[1]


def test_parse_traffic_units() -> None:
    assert parse_traffic_size("1 KB") == 1024
    assert parse_traffic_size("1.5 GB") == round(1.5 * 1024**3)
    assert parse_traffic_size("780 MB") == 780 * 1024**2


def test_update_is_idempotent_and_creates_backup(tmp_path: Path) -> None:
    output = tmp_path / "traffic.xlsx"
    first = update_workbook(ROOT / "data" / "traffic_sample.csv", output)
    second = update_workbook(ROOT / "data" / "traffic_sample.csv", output)
    assert first["inserted"] == 8
    assert second["inserted"] == 0
    assert second["backup"] is not None
    assert Path(second["backup"]).exists()

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames == ["Resumo", "Dados"]
    assert workbook["Dados"].max_row == 9
