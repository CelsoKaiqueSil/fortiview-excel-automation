from pathlib import Path
from unittest.mock import patch
import pandas as pd
import pytest
from openpyxl import load_workbook
from fortiview_excel.pipeline import update_workbook, workbook_lock, backup_workbook

SOURCE = Path(__file__).parents[1] / 'data/traffic_sample.csv'


def test_lock_rejects_simultaneous_run(tmp_path):
    output = tmp_path / 'sample.xlsx'
    with workbook_lock(output):
        with pytest.raises(RuntimeError, match='lock'):
            update_workbook(SOURCE, output)
    assert not output.with_suffix('.xlsx.lock').exists()


def test_failure_preserves_original(tmp_path):
    output = tmp_path / 'sample.xlsx'
    update_workbook(SOURCE, output)
    original = output.read_bytes()
    with patch('fortiview_excel.pipeline.os.replace', side_effect=OSError('Simulated disk failure')):
        with pytest.raises(OSError):
            update_workbook(SOURCE, output)
    assert output.read_bytes() == original
    assert not output.with_suffix('.xlsx.lock').exists()
    assert len(list(tmp_path.glob('*.xlsx'))) == 1


def test_correction_is_audited_without_extra_sample(tmp_path):
    output = tmp_path / 'sample.xlsx'
    update_workbook(SOURCE, output)
    raw = pd.read_csv(SOURCE)
    raw.loc[0, 'value'] = '2 GB'
    changed = tmp_path / 'changed.csv'
    raw.to_csv(changed, index=False)
    result = update_workbook(changed, output)
    assert result['updated'] == 1
    assert result['inserted'] == 0
    book = load_workbook(output)
    try:
        assert book['Dados'].max_row == 9
        assert book['Alteracoes'].max_row == 2
        assert book['Dados'].auto_filter.ref == 'A1:F9'
    finally:
        book.close()


def test_backups_do_not_collide(tmp_path):
    output = tmp_path / 'sample.xlsx'
    output.write_bytes(b'example')
    assert backup_workbook(output) != backup_workbook(output)
